import io
import sqlite3
import pandas as pd
import streamlit as st
from docx import Document
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas

st.set_page_config(
    page_title="SIMANTAP - Manajemen Aset Terpadu",
    page_icon="💎",
    layout="wide",
)

st.markdown("""
<style>
.stApp { background: linear-gradient(135deg,#f8fafc,#eef2ff); }
[data-testid="stSidebar"] { background: linear-gradient(180deg,#0f172a,#1e293b); }
[data-testid="stSidebar"] * { color: #f8fafc !important; }
.hero {
    padding: 24px; border-radius: 18px; color: white;
    background: linear-gradient(135deg,#1e3a8a,#2563eb);
    margin-bottom: 18px;
}
.asset-card {
    padding: 18px; border-radius: 15px; background: white;
    border-left: 6px solid #2563eb; box-shadow: 0 4px 12px #00000012;
    margin-bottom: 12px;
}
.asset-card h4 { margin: 0; color: #1e3a8a; }
.asset-card h2 { margin: 5px 0; color: #0f172a; }
.asset-card p { margin: 0; color: #475569; }
.section {
    padding: 14px 18px; border-radius: 12px; color: white;
    background: linear-gradient(135deg,#0f766e,#059669);
    margin-top: 20px; margin-bottom: 12px;
}
div[data-testid="stMetric"] {
    background: white; padding: 12px; border-radius: 12px;
    box-shadow: 0 3px 10px #0000000d;
}
</style>
""", unsafe_allow_html=True)


@st.cache_data
def load_data(table_name):
    try:
        with sqlite3.connect("simantap_database.db") as conn:
            return pd.read_sql(f"SELECT * FROM {table_name}", conn)
    except Exception:
        return pd.DataFrame()


# Dipertahankan sama seperti script asli agar hasil nilai tidak berubah.
def clean_harga(df, possible_cols=["Harga", "Nilai", "Total_Harga"]):
    target_col = None
    for col in possible_cols:
        if col in df.columns:
            target_col = col
            break
    if target_col is None:
        for c in df.columns:
            if any(kw in c.lower() for kw in ["harga", "nilai", "cost"]):
                target_col = c
                break
    if target_col:
        cleaned = (
            df[target_col].astype(str)
            .str.replace(r"[^\d.]", "", regex=True)
            .replace("", "0")
        )
        return pd.to_numeric(cleaned, errors="coerce").fillna(0)
    return pd.Series(0, index=df.index)


# Nama kategori dipertahankan dari script asli.
def kategorkan_kendaraan(nama_brg):
    if pd.isna(nama_brg):
        return "Lainnya"
    n = str(nama_brg).upper()
    if any(x in n for x in ["ALAT BERAT", "EXCAVATOR", "LOADER", "GRADER", "BULLDOZER"]):
        return "Alat Berat"
    if any(x in n for x in ["PICK", "STRADA", "HILUX", "BOX", "TRUCK"]):
        return "Pick Up / Truk"
    if (
        any(x in n for x in ["MOTOR", "SEPEDA MOTOR", "TRAIL", "VESPA", "YAMAHA", "HONDA"])
        and not any(x in n for x in ["CIVIC", "AVANZA", "INNOVA", "SEDAN"])
    ):
        return "Sepeda Motor"
    return "Mobil Dinas"


def rupiah(value):
    return f"Rp {value:,.0f}"


def detail_downloads(detail_data, prefix, title):
    item_id = str(detail_data.iloc[0]) if len(detail_data) else "ID"
    c1, c2, c3 = st.columns(3)

    out_excel = io.BytesIO()
    with pd.ExcelWriter(out_excel, engine="openpyxl") as writer:
        pd.DataFrame([detail_data]).to_excel(writer, index=False)
    c1.download_button("📊 Excel (.xlsx)", out_excel.getvalue(),
                       file_name=f"{prefix}_{item_id}.xlsx")

    doc = Document()
    doc.add_heading(title, level=1)
    for key, value in detail_data.items():
        doc.add_paragraph(f"{str(key).replace('_', ' ')} : {value}")
    out_word = io.BytesIO()
    doc.save(out_word)
    c2.download_button("📄 Word (.docx)", out_word.getvalue(),
                       file_name=f"{prefix}_{item_id}.docx")

    out_pdf = io.BytesIO()
    pdf = canvas.Canvas(out_pdf, pagesize=letter)
    _, height = letter
    pdf.drawString(40, height - 40, title.upper())
    y = height - 70
    for key, value in detail_data.items():
        pdf.drawString(40, y, f"{str(key).replace('_', ' ')}: {str(value)[:80]}")
        y -= 18
        if y < 40:
            pdf.showPage()
            y = height - 40
    pdf.save()
    c3.download_button("📑 PDF (.pdf)", out_pdf.getvalue(),
                       file_name=f"{prefix}_{item_id}.pdf")


def show_detail(df, selected_key, prefix, title, session_key):
    if selected_key is None:
        return False
    if selected_key not in df.index:
        st.session_state[session_key] = None
        return False

    st.markdown(f'<div class="section">📱 {title}</div>', unsafe_allow_html=True)
    if st.button("⬅️ Kembali ke tabel"):
        st.session_state[session_key] = None
        st.rerun()

    detail = df.loc[selected_key]
    detail_list = [
        {"Atribut / Kolom": str(k).replace("_", " ").title(),
         "Keterangan Detail": str(v if pd.notna(v) else "-")}
        for k, v in detail.items() if k != "Harga_Clean"
    ]
    st.dataframe(pd.DataFrame(detail_list), use_container_width=True,
                 hide_index=True, height=400)
    st.markdown("### 📥 Unduh Dokumen Laporan")
    detail_downloads(detail, prefix, title)
    return True


def show_asset_table(df, title, session_key, prefix, search_label):
    if df.empty:
        st.warning("Data belum tersedia.")
        return

    df = df.copy()
    df["Harga_Clean"] = clean_harga(df)

    if session_key not in st.session_state:
        st.session_state[session_key] = None

    if show_detail(df, st.session_state[session_key], prefix, title, session_key):
        return

    st.markdown(f'<div class="hero"><h2>📋 {title}</h2>'
                f'<p>Pilih satu baris untuk melihat detail dan mengunduh laporan.</p></div>',
                unsafe_allow_html=True)

    st.metric("Jumlah Barang", f"{len(df):,}")
    st.metric("Total Nilai", rupiah(df["Harga_Clean"].sum()))

    keyword = st.text_input(search_label, key=f"search_{session_key}")
    view = df.copy()
    if keyword:
        mask = view.astype(str).apply(
            lambda row: row.str.contains(keyword, case=False, na=False).any(), axis=1
        )
        view = view[mask]

    event = st.dataframe(
        view.drop(columns=["Harga_Clean"], errors="ignore"),
        use_container_width=True, height=420,
        on_select="rerun", selection_mode="single-row",
    )
    if event.selection.rows:
        position = event.selection.rows[0]
        st.session_state[session_key] = view.index[position]
        st.rerun()


def category_summary(df):
    if df.empty:
        return pd.DataFrame(columns=["Kategori", "Jumlah", "Nilai"])
    x = df.copy()
    name_col = (
        "Nama_Barang_Jenis_Barang"
        if "Nama_Barang_Jenis_Barang" in x.columns
        else (x.columns[1] if len(x.columns) > 1 else x.columns[0])
    )
    x["Harga_Clean"] = clean_harga(x)
    x["Kategori"] = x[name_col].apply(kategorkan_kendaraan)
    return x.groupby("Kategori", dropna=False).agg(
        Jumlah=("Kategori", "size"), Nilai=("Harga_Clean", "sum")
    ).reset_index()


def skpd_summary(df, label):
    if df.empty or "SKPD" not in df.columns:
        st.info(f"Belum ada data SKPD untuk {label}.")
        return
    x = df.copy()
    x["Harga_Clean"] = clean_harga(x)
    result = x.groupby("SKPD", dropna=False).agg(
        Jumlah_Barang=("SKPD", "size"),
        Total_Nilai=("Harga_Clean", "sum")
    ).reset_index()
    result["Total_Nilai"] = result["Total_Nilai"].map(rupiah)
    st.dataframe(result, use_container_width=True, hide_index=True)


# HEADER
st.markdown("""
<div class="hero">
<h1>🛡️ BMD - Aset Daerah</h1>
<p>SIMANTAP v2.1 Professional — Sistem Informasi Manajemen Aset Terpadu</p>
</div>
""", unsafe_allow_html=True)

# Semua menu langsung tampil; tidak memakai selectbox "Pilih Modul Aset".
st.sidebar.markdown("## 📌 MENU NAVIGASI")
st.sidebar.markdown("- 🏠 Beranda & Ringkasan")
st.sidebar.markdown("- 🚗 Kendaraan Dinas")
st.sidebar.markdown("- 🗺️ KIB A — Tanah")
st.sidebar.markdown("- 🏢 KIB C — Gedung & Bangunan")

df_k = load_data("tabel_kendaraan")
df_a = load_data("tabel_kib_a_tanah")
df_c = load_data("tabel_kib_c_gedung")

# Ringkasan kendaraan berdasarkan kategori.
st.markdown('<div class="section">🏠 BERANDA & RINGKASAN EKSEKUTIF</div>',
            unsafe_allow_html=True)
summary = category_summary(df_k)
category_order = ["Alat Berat", "Mobil Dinas", "Pick Up / Truk", "Sepeda Motor"]
cards = st.columns(4)
for col, category in zip(cards, category_order):
    row = summary[summary["Kategori"] == category]
    count = int(row["Jumlah"].iloc[0]) if not row.empty else 0
    value = float(row["Nilai"].iloc[0]) if not row.empty else 0
    icon = {"Alat Berat": "🚜", "Mobil Dinas": "🚗",
            "Pick Up / Truk": "🚚", "Sepeda Motor": "🏍️"}[category]
    with col:
        st.markdown(f"""<div class="asset-card">
        <h4>{icon} {category}</h4><h2>{count:,} Unit</h2>
        <p>{rupiah(value)}</p></div>""", unsafe_allow_html=True)

st.markdown("### 📊 Jumlah dan Nilai Kendaraan per Kategori")
if not summary.empty:
    display_summary = summary.copy()
    display_summary["Nilai"] = display_summary["Nilai"].map(rupiah)
    st.dataframe(display_summary, use_container_width=True, hide_index=True)
    chart = summary.set_index("Kategori")[["Jumlah"]]
    st.bar_chart(chart)

st.markdown("### 🏢 Jumlah Barang dan Nilai per SKPD — Kendaraan")
skpd_summary(df_k, "Kendaraan")

st.markdown("### 🗺️ Jumlah Barang dan Nilai per SKPD — KIB A Tanah")
skpd_summary(df_a, "KIB A")

st.markdown("### 🏢 Jumlah Barang dan Nilai per SKPD — KIB C Gedung")
skpd_summary(df_c, "KIB C")

# Semua tabel tampil berurutan.
st.markdown('<div class="section">🚗 KENDARAAN DINAS</div>', unsafe_allow_html=True)
show_asset_table(df_k, "Manajemen Aset Kendaraan Dinas",
                 "selected_vehicle", "Kendaraan",
                 "🔎 Cari No. Polisi / Merk / Pengguna")

st.markdown('<div class="section">🗺️ KIB A — TANAH</div>', unsafe_allow_html=True)
show_asset_table(df_a, "Manajemen Aset KIB A (Tanah)",
                 "selected_land", "Tanah",
                 "🔎 Cari Alamat / Lokasi")

st.markdown('<div class="section">🏢 KIB C — GEDUNG & BANGUNAN</div>', unsafe_allow_html=True)
show_asset_table(df_c, "Manajemen Aset KIB C (Gedung & Bangunan)",
                 "selected_building", "Gedung",
                 "🔎 Cari Nama Gedung / Lokasi")
